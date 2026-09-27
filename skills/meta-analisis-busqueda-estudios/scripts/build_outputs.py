#!/usr/bin/env python3
"""Build the screening deliverables from <workdir>/pool.json + decisions + search log.

Outputs:
  <out_prefix>.xlsx  sheets: Candidates (High+Medium), Reviews & meta-analyses, Low & excluded, Search log
  <out_prefix>.ris   High + Medium + Reviews, built from stored metadata (no Scite call needed)

decisions.json: { "<doi>": {"tier": "H|M|R|L|X", "system": "...", "comparison": "...",
                            "data": "...", "why": "..."} }
  H = high probability, M = medium, R = review/meta-analysis, L = low, X = excluded.
  Pool records without a decision are written to "Low & excluded" as off-topic.
search_log.json: {"Field": "text", ...} written verbatim to the Search log sheet (insertion order kept).

Usage: build_outputs.py <workdir> <decisions.json> <search_log.json> <out_prefix> [--focus '<regex>']
  --focus: within each tier, rows whose 'system' matches the regex are listed first
           (e.g. the primary taxon/outcome when the question has a primary and a secondary scope).
"""
import json, os, re, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

args = sys.argv[1:]
focus = None
if '--focus' in args:
    i = args.index('--focus'); focus = re.compile(args[i + 1], re.I); args = args[:i] + args[i + 2:]
workdir, dec_path, log_path, out = args
pool = json.load(open(os.path.join(workdir, 'pool.json')))
dec = {k.lower(): v for k, v in json.load(open(dec_path)).items()}
log = json.load(open(log_path))

TIER = {'H': 'High', 'M': 'Medium', 'R': 'Review/Meta-analysis', 'L': 'Low', 'X': 'Excluded'}
ORDER = {t: n for n, t in enumerate('HMRLX')}
COLOR = {'H': 'C6EFCE', 'M': 'FFEB9C', 'R': 'DDEBF7', 'L': 'F2F2F2', 'X': 'F2F2F2'}

rows = []
for doi, v in pool.items():
    d = dec.get(doi) or {'tier': 'X', 'why': 'Off-topic by title/abstract'}
    rows.append({**v, 'tier': d.get('tier', 'X'), 'system': d.get('system', ''), 'comparison': d.get('comparison', ''),
                 'data': d.get('data', ''), 'why': d.get('why', '')})
unknown = set(dec) - set(pool)
if unknown:
    print('WARNING: decisions for DOIs not in pool (ignored):', ', '.join(sorted(unknown)))


def first_author(r):
    a = r.get('authors') or []
    return (a[0] if a else '') + (' et al.' if len(a) > 1 else '')


def abstract_src(r):
    return r.get('abstract_src') or 'scite' if r.get('abstract') and r.get('abstract_src') != 'none' else 'none (title only)'


BASE = [('#', lambda r, n: n, 5), ('Probability', lambda r, n: TIER[r['tier']], 13),
        ('First author', lambda r, n: first_author(r), 20), ('Year', lambda r, n: r.get('year'), 7),
        ('Title', lambda r, n: r['title'], 60), ('Journal', lambda r, n: r.get('journal'), 24),
        ('DOI', lambda r, n: r['doi'], 28)]
CAND = BASE + [('Study system', lambda r, n: r['system'], 24),
               ('Comparison / predictor', lambda r, n: r['comparison'], 40),
               ('Expected data', lambda r, n: r['data'], 32),
               ('Why probable / caveats', lambda r, n: r['why'], 48),
               ('Abstract source', lambda r, n: abstract_src(r), 16),
               ('Found by', lambda r, n: ', '.join(r.get('found_by', [])), 16),
               ('Your decision', lambda r, n: '', 14), ('Notes', lambda r, n: '', 30)]
hdr_font, hdr_fill = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='305496')


def write_sheet(ws, sel, cols):
    ws.append([c[0] for c in cols])
    for c in ws[1]:
        c.font, c.fill = hdr_font, hdr_fill
        c.alignment = Alignment(wrap_text=True, vertical='center')
    sel = sorted(sel, key=lambda r: (ORDER[r['tier']], 0 if focus and focus.search(r['system']) else 1,
                                     -(r.get('year') or 0)))
    doi_col = [h for h, _, _ in cols].index('DOI') + 1
    for n, r in enumerate(sel, 1):
        ws.append([f(r, n) for _, f, _ in cols])
        row = ws.max_row
        for c in ws[row]:
            c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.cell(row, 2).fill = PatternFill('solid', fgColor=COLOR[r['tier']])
        cell = ws.cell(row, doi_col)
        cell.hyperlink = 'https://doi.org/' + r['doi']; cell.font = Font(color='0563C1', underline='single')
    for j, (_, _, w) in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(1, j).column_letter].width = w
    ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions
    return len(sel)


wb = Workbook()
ws = wb.active; ws.title = 'Candidates'
n_c = write_sheet(ws, [r for r in rows if r['tier'] in 'HM'], CAND)
if n_c:
    dv = DataValidation(type='list', formula1='"Include,Exclude,Maybe,No access"', allow_blank=True)
    ws.add_data_validation(dv)
    col = ws.cell(1, [h for h, _, _ in CAND].index('Your decision') + 1).column_letter
    dv.add(f'{col}2:{col}{ws.max_row}')
n_r = write_sheet(wb.create_sheet('Reviews & meta-analyses'), [r for r in rows if r['tier'] == 'R'],
                  BASE + [('Scope', lambda r, n: r['comparison'], 45), ('Use', lambda r, n: r['why'], 50)])
n_x = write_sheet(wb.create_sheet('Low & excluded'), [r for r in rows if r['tier'] in 'LX'],
                  BASE + [('Reason', lambda r, n: r['why'], 55)])
lw = wb.create_sheet('Search log')
lw.column_dimensions['A'].width, lw.column_dimensions['B'].width = 28, 110
cnt = {t: sum(r['tier'] == t for r in rows) for t in 'HMRLX'}
log = dict(log, **{'Counts': f"Screened {len(rows)} records: High {cnt['H']}, Medium {cnt['M']}, "
                             f"Reviews/meta-analyses {cnt['R']}, Low {cnt['L']}, Excluded {cnt['X']}"})
for k, v in log.items():
    lw.append([k, v if isinstance(v, str) else json.dumps(v, ensure_ascii=False)])
    lw.cell(lw.max_row, 1).font = Font(bold=True)
    lw.cell(lw.max_row, 2).alignment = Alignment(wrap_text=True, vertical='top')
wb.save(out + '.xlsx')

with open(out + '.ris', 'w', encoding='utf-8') as f:
    for r in sorted([r for r in rows if r['tier'] in 'HMR'], key=lambda r: ORDER[r['tier']]):
        f.write('TY  - JOUR\n')
        for a in r.get('authors') or []:
            f.write(f'AU  - {a}\n')
        f.write(f"TI  - {r['title']}\nPY  - {r.get('year') or ''}\nJO  - {r.get('journal') or ''}\n")
        for tag, key in (('VL', 'volume'), ('IS', 'issue'), ('SP', 'page')):
            if r.get(key):
                f.write(f'{tag}  - {r[key]}\n')
        f.write(f"DO  - {r['doi']}\nUR  - https://doi.org/{r['doi']}\nKW  - {TIER[r['tier']]}\n")
        if r['why']:
            f.write(f"N1  - {r['why']}\n")
        if r.get('abstract') and r.get('abstract_src') != 'none':
            f.write(f"AB  - {r['abstract']}\n")
        f.write('ER  - \n\n')
print(f"counts={cnt} candidates={n_c} reviews={n_r} low_excluded={n_x}\n-> {out}.xlsx\n-> {out}.ris")
